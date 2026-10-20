"""Protótipo DZOO: gestão de demandas, atas, aquisições e recebimentos."""

import io
import sqlite3
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import streamlit as st
from docx import Document
from importador import baixar_drive_publico, criar_previa

st.set_page_config(page_title="DZOO | Gestão de Aquisições", page_icon="📦", layout="wide")
DB = Path(__file__).with_name("dzoo_demo.db")


def conn():
    connection = sqlite3.connect(DB)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def sql(query, params=()):
    with conn() as c:
        return pd.read_sql_query(query, c, params=params)


def one(query, params=()):
    with conn() as c:
        return c.execute(query, params).fetchone()


def execute(query, params=()):
    with conn() as c:
        cur = c.execute(query, params)
        c.commit()
        return cur.lastrowid


def init():
    statements = [
        "CREATE TABLE IF NOT EXISTS setores (id INTEGER PRIMARY KEY, nome TEXT UNIQUE NOT NULL, demo INT DEFAULT 0)",
        "CREATE TABLE IF NOT EXISTS produtos (id INTEGER PRIMARY KEY, nome TEXT NOT NULL, codigo_sipac TEXT, unidade TEXT NOT NULL, grupo TEXT, especificacao TEXT, demo INT DEFAULT 0)",
        "CREATE TABLE IF NOT EXISTS demandas (id INTEGER PRIMARY KEY, ano INT NOT NULL, setor_id INT NOT NULL, produto_id INT NOT NULL, quantidade REAL NOT NULL CHECK(quantidade>=0), confirmado INT DEFAULT 0, demo INT DEFAULT 0, FOREIGN KEY(setor_id) REFERENCES setores, FOREIGN KEY(produto_id) REFERENCES produtos)",
        "CREATE TABLE IF NOT EXISTS atas (id INTEGER PRIMARY KEY, numero TEXT NOT NULL UNIQUE, pregao TEXT, fornecedor TEXT NOT NULL, email TEXT, telefone TEXT, inicio TEXT, fim TEXT, url TEXT, interesse INT DEFAULT 1, demo INT DEFAULT 0)",
        "CREATE TABLE IF NOT EXISTS ata_itens (id INTEGER PRIMARY KEY, ata_id INT NOT NULL, produto_id INT NOT NULL, item_num TEXT NOT NULL, quantidade REAL NOT NULL CHECK(quantidade>=0), preco REAL NOT NULL CHECK(preco>=0), unidade TEXT NOT NULL, demo INT DEFAULT 0, UNIQUE(ata_id,item_num), FOREIGN KEY(ata_id) REFERENCES atas, FOREIGN KEY(produto_id) REFERENCES produtos)",
        "CREATE TABLE IF NOT EXISTS pplans (id INTEGER PRIMARY KEY, numero TEXT, data TEXT NOT NULL, situacao TEXT NOT NULL, observacoes TEXT, demo INT DEFAULT 0)",
        "CREATE TABLE IF NOT EXISTS pplan_itens (id INTEGER PRIMARY KEY, pplan_id INT NOT NULL, ata_item_id INT NOT NULL, quantidade REAL NOT NULL CHECK(quantidade>0), demo INT DEFAULT 0, FOREIGN KEY(pplan_id) REFERENCES pplans, FOREIGN KEY(ata_item_id) REFERENCES ata_itens)",
        "CREATE TABLE IF NOT EXISTS autorizacoes (id INTEGER PRIMARY KEY, pplan_item_id INT NOT NULL, data TEXT NOT NULL, quantidade REAL NOT NULL CHECK(quantidade>0), referencia TEXT, demo INT DEFAULT 0, FOREIGN KEY(pplan_item_id) REFERENCES pplan_itens)",
        "CREATE TABLE IF NOT EXISTS requisicoes (id INTEGER PRIMARY KEY, numero TEXT UNIQUE NOT NULL, data TEXT NOT NULL, observacoes TEXT, demo INT DEFAULT 0)",
        "CREATE TABLE IF NOT EXISTS req_itens (id INTEGER PRIMARY KEY, requisicao_id INT NOT NULL, ata_item_id INT NOT NULL, pplan_item_id INT, quantidade REAL NOT NULL CHECK(quantidade>0), demo INT DEFAULT 0, FOREIGN KEY(requisicao_id) REFERENCES requisicoes, FOREIGN KEY(ata_item_id) REFERENCES ata_itens, FOREIGN KEY(pplan_item_id) REFERENCES pplan_itens)",
        "CREATE TABLE IF NOT EXISTS empenhos (id INTEGER PRIMARY KEY, numero TEXT UNIQUE NOT NULL, requisicao_id INT, data TEXT, status TEXT DEFAULT 'Aberto', demo INT DEFAULT 0, FOREIGN KEY(requisicao_id) REFERENCES requisicoes)",
        "CREATE TABLE IF NOT EXISTS empenho_itens (id INTEGER PRIMARY KEY, empenho_id INT NOT NULL, ata_item_id INT NOT NULL, quantidade REAL NOT NULL CHECK(quantidade>0), preco REAL NOT NULL CHECK(preco>=0), demo INT DEFAULT 0, FOREIGN KEY(empenho_id) REFERENCES empenhos, FOREIGN KEY(ata_item_id) REFERENCES ata_itens)",
        "CREATE TABLE IF NOT EXISTS entregas (id INTEGER PRIMARY KEY, empenho_id INT NOT NULL, data TEXT NOT NULL, local TEXT, conferido INT DEFAULT 0, observacoes TEXT, demo INT DEFAULT 0, FOREIGN KEY(empenho_id) REFERENCES empenhos)",
        "CREATE TABLE IF NOT EXISTS notas (id INTEGER PRIMARY KEY, entrega_id INT NOT NULL, numero TEXT NOT NULL, serie TEXT, valor REAL, enviado INT DEFAULT 0, data_envio TEXT, demo INT DEFAULT 0, FOREIGN KEY(entrega_id) REFERENCES entregas)",
        "CREATE TABLE IF NOT EXISTS nota_itens (id INTEGER PRIMARY KEY, nota_id INT NOT NULL, empenho_item_id INT NOT NULL, quantidade REAL NOT NULL CHECK(quantidade>0), demo INT DEFAULT 0, FOREIGN KEY(nota_id) REFERENCES notas, FOREIGN KEY(empenho_item_id) REFERENCES empenho_itens)",
    ]
    with conn() as c:
        for statement in statements:
            c.execute(statement)
        c.commit()


def seed():
    if one("SELECT COUNT(*) AS n FROM produtos")["n"] or one("SELECT COUNT(*) AS n FROM atas")["n"]:
        st.warning("O banco já possui registros. A carga fictícia só funciona em uma base vazia.")
        return
    setor = execute("INSERT INTO setores (nome,demo) VALUES (?,1)", ("Caprinocultura",))
    for s in ["Apicultura", "Ovinocultura", "Coturnicultura", "Cunicultura", "Equideocultura", "Helicicultura", "Forragicultura", "Aquacultura"]:
        execute("INSERT INTO setores (nome,demo) VALUES (?,1)", (s,))
    p1 = execute("INSERT INTO produtos (nome,codigo_sipac,unidade,grupo,demo) VALUES (?,?,?,?,1)", ("Fubá de milho", "EXEMPLO-001", "kg", "3006"))
    p2 = execute("INSERT INTO produtos (nome,codigo_sipac,unidade,grupo,demo) VALUES (?,?,?,?,1)", ("Farelo de soja", "EXEMPLO-002", "kg", "3006"))
    p3 = execute("INSERT INTO produtos (nome,codigo_sipac,unidade,grupo,demo) VALUES (?,?,?,?,1)", ("Ração para coelhos", "EXEMPLO-003", "kg", "3006"))
    a = execute("INSERT INTO atas (numero,pregao,fornecedor,email,telefone,inicio,fim,demo) VALUES (?,?,?,?,?,?,?,1)", ("DEMO-01/2026", "Pregão fictício", "Fornecedor Demonstração Ltda.", "contato@exemplo.invalid", "(00) 0000-0000", "2026-01-01", "2027-01-01"))
    i1=execute("INSERT INTO ata_itens (ata_id,produto_id,item_num,quantidade,preco,unidade,demo) VALUES (?,?,?,?,?,?,1)",(a,p1,"1",5000,2.63,"kg"))
    i2=execute("INSERT INTO ata_itens (ata_id,produto_id,item_num,quantidade,preco,unidade,demo) VALUES (?,?,?,?,?,?,1)",(a,p2,"2",2500,3.4,"kg"))
    execute("INSERT INTO ata_itens (ata_id,produto_id,item_num,quantidade,preco,unidade,demo) VALUES (?,?,?,?,?,?,1)",(a,p3,"3",800,4.14,"kg"))
    execute("INSERT INTO demandas (ano,setor_id,produto_id,quantidade,confirmado,demo) VALUES (?,?,?,?,?,1)",(2027,setor,p1,1000,1))
    pp=execute("INSERT INTO pplans (numero,data,situacao,demo) VALUES (?,?,?,1)",("MEM-DEMO/2026","2026-10-01","Parcial"))
    pi=execute("INSERT INTO pplan_itens (pplan_id,ata_item_id,quantidade,demo) VALUES (?,?,?,1)",(pp,i1,1000))
    execute("INSERT INTO autorizacoes (pplan_item_id,data,quantidade,referencia,demo) VALUES (?,?,?,?,1)",(pi,"2026-10-02",500,"Autorização fictícia"))
    req=execute("INSERT INTO requisicoes (numero,data,demo) VALUES (?,?,1)",("REQ-DEMO/2026","2026-10-03"))
    execute("INSERT INTO req_itens (requisicao_id,ata_item_id,pplan_item_id,quantidade,demo) VALUES (?,?,?,?,1)",(req,i1,pi,500))
    emp=execute("INSERT INTO empenhos (numero,requisicao_id,data,demo) VALUES (?,?,?,1)",("EMP-DEMO/2026",req,"2026-10-04"))
    ei=execute("INSERT INTO empenho_itens (empenho_id,ata_item_id,quantidade,preco,demo) VALUES (?,?,?,?,1)",(emp,i1,500,2.63))
    en=execute("INSERT INTO entregas (empenho_id,data,local,conferido,demo) VALUES (?,?,?,?,1)",(emp,"2026-10-07","CTAN",1))
    n1=execute("INSERT INTO notas (entrega_id,numero,serie,valor,enviado,data_envio,demo) VALUES (?,?,?,?,?,?,1)",(en,"DEMO-4407","1",263,1,"2026-10-08"))
    n2=execute("INSERT INTO notas (entrega_id,numero,serie,valor,enviado,demo) VALUES (?,?,?,?,?,1)",(en,"DEMO-4408","1",263,0))
    for n in [n1,n2]:
        execute("INSERT INTO nota_itens (nota_id,empenho_item_id,quantidade,demo) VALUES (?,?,?,1)",(n,ei,100))
    st.success("Dados fictícios criados. Nenhum valor corresponde à situação oficial da UFSJ.")


def wipe_demo():
    names=["nota_itens","notas","entregas","empenho_itens","empenhos","req_itens","requisicoes","autorizacoes","pplan_itens","pplans","demandas","ata_itens","atas","produtos","setores"]
    with conn() as c:
        for name in names:
            c.execute(f"DELETE FROM {name} WHERE demo=1")
        c.commit()


def options(df, display):
    return {f"{str(display(row))}  [ID {int(row['id'])}]": int(row["id"]) for _,row in df.iterrows()}


def pick(label, df, display, key=None):
    opts=options(df,display)
    if not opts:
        st.info(f"Cadastre os registros necessários para: {label.lower()}.")
        return None
    chosen=st.selectbox(label,list(opts),key=key)
    return opts[chosen]


def money(value):
    return f"R$ {float(value or 0):,.2f}".replace(",","X").replace(".",",").replace("X",".")


def table(df):
    if df.empty:
        st.info("Nenhum registro encontrado.")
    else:
        st.dataframe(df, hide_index=True, use_container_width=True)


def download_xlsx(df, name):
    buf=io.BytesIO()
    with pd.ExcelWriter(buf,engine="openpyxl") as writer:
        df.to_excel(writer,index=False,sheet_name="Relatorio")
    st.download_button("⬇️ Exportar Excel",buf.getvalue(),name,mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",disabled=df.empty)


def view_atas():
    return sql("""SELECT ai.id,p.nome AS produto,p.codigo_sipac,ai.item_num,ai.unidade,ai.quantidade,ai.preco, a.numero AS ata,a.fornecedor,a.email,a.telefone,a.fim, a.id AS ata_id,
     COALESCE((SELECT SUM(ei.quantidade) FROM empenho_itens ei WHERE ei.ata_item_id=ai.id),0) AS empenhado,
     COALESCE((SELECT SUM(ri.quantidade) FROM req_itens ri WHERE ri.ata_item_id=ai.id AND NOT EXISTS(SELECT 1 FROM empenhos e WHERE e.requisicao_id=ri.requisicao_id)),0) AS requisitado_sem_empenho
     FROM ata_itens ai JOIN atas a ON a.id=ai.ata_id JOIN produtos p ON p.id=ai.produto_id WHERE a.interesse=1 ORDER BY p.nome,a.numero""")


def pendencias():
    return sql("""SELECT e.numero AS empenho,r.numero AS requisicao,a.numero AS ata,p.nome AS produto,a.fornecedor,a.email,a.telefone,ei.id AS empenho_item_id,ei.quantidade AS empenhado,
      COALESCE((SELECT SUM(ni.quantidade) FROM nota_itens ni JOIN notas n ON n.id=ni.nota_id JOIN entregas ent ON ent.id=n.entrega_id WHERE ni.empenho_item_id=ei.id AND ent.conferido=1),0) AS recebido,
      ei.preco,ei.quantidade - COALESCE((SELECT SUM(ni.quantidade) FROM nota_itens ni JOIN notas n ON n.id=ni.nota_id JOIN entregas ent ON ent.id=n.entrega_id WHERE ni.empenho_item_id=ei.id AND ent.conferido=1),0) AS pendente
      FROM empenho_itens ei JOIN empenhos e ON e.id=ei.empenho_id LEFT JOIN requisicoes r ON r.id=e.requisicao_id JOIN ata_itens ai ON ai.id=ei.ata_item_id JOIN atas a ON a.id=ai.ata_id JOIN produtos p ON p.id=ai.produto_id WHERE e.status!='Cancelado' ORDER BY e.numero,p.nome""")


def home():
    st.title("📦 DZOO | Gestão de Aquisições")
    st.caption("Protótipo v0.2 · Registro local de testes · Informações não integradas ao SIPAC")
    p=pendencias(); a=view_atas()
    n_notas=one("SELECT COUNT(*) AS n FROM notas WHERE enviado=0")["n"]
    x,y,z,w=st.columns(4)
    x.metric("Produtos cadastrados",int(one("SELECT COUNT(*) AS n FROM produtos")["n"]))
    y.metric("Atas de interesse",int(one("SELECT COUNT(*) AS n FROM atas WHERE interesse=1")["n"]))
    z.metric("Itens de empenho pendentes",int((p.pendente>0.000001).sum()) if not p.empty else 0)
    w.metric("DANFEs a encaminhar",int(n_notas))
    st.subheader("Acompanhamento prioritário")
    c1,c2=st.columns(2)
    with c1:
        st.markdown("**Entregas pendentes**")
        table(p[p.pendente>0][["produto","fornecedor","empenho","pendente"]] if not p.empty else p)
    with c2:
        st.markdown("**Itens com saldo calculado**")
        if not a.empty:
            a["saldo_calculado"]=(a.quantidade-a.empenhado).clip(lower=0)
            table(a[a.saldo_calculado>0][["produto","ata","saldo_calculado","unidade"]])
        else: st.info("Sem atas cadastradas.")
    st.warning("Saldos calculados apenas com empenhos cadastrados nesta base. **Não representam o saldo oficial SIPAC**.")


def produtos():
    st.title("Produtos e setores")
    t1,t2,t3=st.tabs(["Consultar produtos","Novo produto","Setores"])
    with t1:
        termo=st.text_input("Buscar produto por nome/código")
        df=sql("SELECT id,nome,codigo_sipac,unidade,grupo,especificacao FROM produtos ORDER BY nome")
        if termo: df=df[df.astype(str).apply(lambda col: col.str.contains(termo,case=False,na=False)).any(axis=1)]
        table(df)
        if not df.empty:
            selected=pick("Ficha do produto",df,lambda r:r["nome"])
            if selected:
                st.markdown("**Atas relacionadas**")
                a=view_atas(); items=sql("SELECT id FROM ata_itens WHERE produto_id=?",(selected,)).id.tolist()
                table(a[a.id.isin(items)][["ata","item_num","fornecedor","quantidade","unidade","preco"]] if not a.empty else a)
                st.markdown("**Demandas por setor**")
                table(sql("SELECT d.ano,s.nome AS setor,d.quantidade,p.unidade,d.confirmado FROM demandas d JOIN setores s ON s.id=d.setor_id JOIN produtos p ON p.id=d.produto_id WHERE d.produto_id=? ORDER BY d.ano DESC",(selected,)))
    with t2:
        with st.form("novo_produto"):
            nome=st.text_input("Nome resumido *")
            code=st.text_input("Código SIPAC (opcional)")
            un=st.text_input("Unidade oficial *",value="kg")
            grupo=st.text_input("Grupo (ex.: 3006)")
            spec=st.text_area("Especificação")
            if st.form_submit_button("Cadastrar produto"):
                if nome.strip() and un.strip():
                    execute("INSERT INTO produtos (nome,codigo_sipac,unidade,grupo,especificacao) VALUES (?,?,?,?,?)",(nome.strip(),code.strip(),un.strip(),grupo,spec))
                    st.success("Produto cadastrado.")
                else: st.error("Informe nome e unidade.")
    with t3:
        table(sql("SELECT id,nome FROM setores ORDER BY nome"))
        with st.form("setores_form"):
            name=st.text_input("Novo setor")
            if st.form_submit_button("Adicionar setor"):
                try:
                    execute("INSERT INTO setores (nome) VALUES (?)",(name.strip(),))
                    st.success("Setor cadastrado.")
                except sqlite3.IntegrityError: st.error("Setor vazio ou já cadastrado.")


def planejamento():
    st.title("Planejamento anual")
    st.caption("Demandas dos setores não são resultados de licitação. A referência oficial futura deve ser a quantidade lançada no SIPAC.")
    a,b,c=st.tabs(["Edição por setor","Visão consolidada","Referência SIPAC"])
    with a:
        setores=sql("SELECT * FROM setores ORDER BY nome"); produtos_df=sql("SELECT * FROM produtos ORDER BY nome")
        ano=st.number_input("Ano do planejamento",2020,2100,2027,key="plan_ano")
        sid=pick("Setor",setores,lambda r:r["nome"])
        if sid:
            table(sql("SELECT d.id,p.nome AS produto,d.quantidade,p.unidade,d.confirmado FROM demandas d JOIN produtos p ON p.id=d.produto_id WHERE d.ano=? AND d.setor_id=?",(ano,sid)))
            with st.form("demanda"):
                pid=pick("Produto",produtos_df,lambda r:r["nome"])
                qtd=st.number_input("Quantidade revisada",min_value=0.0,step=1.0)
                conf=st.checkbox("Demanda confirmada pelo setor")
                if st.form_submit_button("Registrar demanda") and pid:
                    execute("INSERT INTO demandas (ano,setor_id,produto_id,quantidade,confirmado) VALUES (?,?,?,?,?)",(ano,sid,pid,qtd,int(conf)))
                    st.success("Demanda registrada; histórico preservado.")
    with b:
        ano2=st.number_input("Ano para consolidação",2020,2100,2027,key="plan_ano2")
        df=sql("SELECT p.nome AS produto,p.unidade,s.nome AS setor,d.quantidade FROM demandas d JOIN produtos p ON p.id=d.produto_id JOIN setores s ON s.id=d.setor_id WHERE d.ano=?",(ano2,))
        if not df.empty:
            pivot=df.pivot_table(index=["produto","unidade"],columns="setor",values="quantidade",aggfunc="sum",fill_value=0)
            pivot["TOTAL"]=pivot.sum(axis=1)
            table(pivot.reset_index()); download_xlsx(pivot.reset_index(),f"planejamento_dzoo_{ano2}.xlsx")
        else: st.info("Sem demandas para esse exercício.")
    with c:
        st.info("Registrar a quantidade consolidada efetivamente lançada no SIPAC na área Requisições. Resultado deserto ou fracassado não zera a demanda histórica.")
        table(sql("SELECT r.numero AS requisicao, p.nome AS produto, SUM(ri.quantidade) AS quantidade_sipac,ai.unidade FROM req_itens ri JOIN requisicoes r ON r.id=ri.requisicao_id JOIN ata_itens ai ON ai.id=ri.ata_item_id JOIN produtos p ON p.id=ai.produto_id GROUP BY r.numero,p.nome,ai.unidade"))


def salvar_itens_selecionados(df):
    """Persiste apenas itens marcados, em transação atômica e sem duplicidade."""
    selected=df[df["Selecionar"] == True].copy()
    if selected.empty:
        raise ValueError("Selecione ao menos um item na prévia.")
    if selected[["Pregão","Ata","Fornecedor","Item","Produto / especificação","Unidade"]].isna().any().any():
        raise ValueError("Há campos vazios nos itens selecionados.")
    count=0; duplicates=0
    with conn() as c:
        for _,r in selected.iterrows():
            pregao=str(r["Pregão"]).strip(); ata=str(r["Ata"]).strip()
            fornecedor=str(r["Fornecedor"]).strip(); num=str(r["Item"]).strip()
            nome=str(r["Produto / especificação"]).strip();unidade=str(r["Unidade"]).strip()
            qtd=float(r["Quantidade"]);preco=float(r["Preço unitário"])
            if not all([pregao,ata,fornecedor,num,nome,unidade]) or qtd<=0 or preco<0:
                raise ValueError("Revise pregão, ata, fornecedor, item, descrição, quantidade e preço dos selecionados.")
            exists=c.execute("SELECT id FROM atas WHERE numero=?",(ata,)).fetchone()
            if exists:
                aid=exists["id"]
                prev=c.execute("SELECT pregao FROM atas WHERE id=?",(aid,)).fetchone()
                if prev["pregao"] and prev["pregao"]!=pregao:
                    raise ValueError(f"Ata {ata} já pertence a outro pregão: {prev['pregao']}.")
            else:
                cur=c.execute("INSERT INTO atas(numero,pregao,fornecedor,email,telefone,url) VALUES(?,?,?,?,?,?)",(ata,pregao,fornecedor,str(r.get("E-mail", "")),str(r.get("Telefone", "")),str(r["Documento"])))
                aid=cur.lastrowid
            if c.execute("SELECT id FROM ata_itens WHERE ata_id=? AND item_num=?",(aid,num)).fetchone():
                duplicates+=1;continue
            # Igualdade exata de descrição e unidade; sem associação automática por semelhança textual.
            p=c.execute("SELECT id FROM produtos WHERE nome=? AND unidade=?",(nome,unidade)).fetchone()
            if p: pid=p["id"]
            else:
                pid=c.execute("INSERT INTO produtos(nome,unidade,especificacao) VALUES(?,?,?)",(nome,unidade,nome)).lastrowid
            c.execute("INSERT INTO ata_itens(ata_id,produto_id,item_num,quantidade,preco,unidade) VALUES(?,?,?,?,?,?)",(aid,pid,num,qtd,preco,unidade))
            count+=1
        c.commit()
    return count,duplicates


def licitacoes():
    st.title("📑 Pregões e itens de interesse do DEZOO")
    st.caption("Organização prioritária por pregão. Apenas itens escolhidos na importação são cadastrados.")
    tabs=st.tabs(["Catálogo por pregão","Importar e selecionar itens","Cadastrar pregão/ata manualmente","Cadastrar item manualmente"])
    with tabs[0]:
        df=view_atas()
        if df.empty: st.info("Nenhum item cadastrado. Importe PDFs ou cadastre itens manualmente.")
        else:
            atas=sql("SELECT id,numero,pregao FROM atas")
            df=df.merge(atas[["id","pregao"]],left_on="ata_id",right_on="id",suffixes=("","_ata"))
            pregoes=sorted(df.pregao.fillna("Não identificado").unique().tolist())
            chosen=st.selectbox("Pregão",pregoes)
            df=df[df.pregao.fillna("Não identificado")==chosen].copy()
            query=st.text_input("Pesquisar produto, item ou fornecedor",key="consulta_pregao")
            if query:
                mask=df[["produto","item_num","fornecedor"]].astype(str).apply(lambda c:c.str.contains(query,case=False,regex=False)).any(axis=1)
                df=df[mask]
            if not df.empty:
                df["saldo_calculado"]=df.quantidade-df.empenhado
                table(df[["item_num","produto","fornecedor","ata","quantidade","unidade","preco","empenhado","saldo_calculado","requisitado_sem_empenho"]].rename(columns={"item_num":"Item","produto":"Produto","ata":"Ata de origem"}))
                download_xlsx(df,"dzoo_itens_pregao.xlsx")
            st.caption("Saldo estimado apenas com empenhos registrados neste aplicativo; não é saldo oficial SIPAC.")
    with tabs[1]:
        st.markdown("**1. Obter arquivos PDF**")
        mode=st.radio("Fonte",["Anexar PDFs","Link público Google Drive"],horizontal=True)
        if mode=="Anexar PDFs":
            files=st.file_uploader("Selecione um ou vários PDFs",type="pdf",accept_multiple_files=True)
            if st.button("Ler PDFs",disabled=not files):
                st.session_state["importacao_arquivos"]=[(f.name,f.getvalue()) for f in files]
        else:
            url=st.text_input("Link público de pasta ou arquivo",placeholder="https://drive.google.com/drive/folders/...")
            if st.button("Buscar PDFs do Drive",disabled=not url.strip()):
                try:
                    with st.spinner("Baixando PDFs públicos..."):
                        st.session_state["importacao_arquivos"]=baixar_drive_publico(url.strip())
                    st.success(f"Encontrados {len(st.session_state['importacao_arquivos'])} PDFs.")
                except Exception as exc: st.error(f"Não foi possível obter os arquivos: {exc}")
        if st.session_state.get("importacao_arquivos"):
            if st.button("Extrair itens e preparar conferência"):
                with st.spinner("Analisando tabelas dos documentos..."):
                    preview,alerts=criar_previa(st.session_state["importacao_arquivos"])
                st.session_state["importacao_previa"]=preview
                st.session_state["importacao_alertas"]=alerts
        if "importacao_previa" in st.session_state:
            st.markdown("**2. Conferir os dados e marcar os itens de interesse**")
            for warning in st.session_state.get("importacao_alertas",[]): st.warning(warning)
            original=st.session_state["importacao_previa"]
            if original.empty: st.error("Nenhum item foi extraído com segurança. Tente PDFs com texto selecionável ou cadastre os itens manualmente.")
            else:
                st.info(f"{len(original)} itens encontrados. Nenhum será salvo sem sua confirmação.")
                editor=st.data_editor(original,use_container_width=True,hide_index=True,num_rows="fixed",key="editor_itens_importados",column_config={"Selecionar":st.column_config.CheckboxColumn("Importar para DEZOO",default=False),"Quantidade":st.column_config.NumberColumn(min_value=0),"Preço unitário":st.column_config.NumberColumn(min_value=0)})
                chosen=editor[editor["Selecionar"]==True]
                st.metric("Itens marcados para cadastro",len(chosen))
                if st.button("Confirmar e salvar somente selecionados",type="primary",disabled=chosen.empty):
                    try:
                        n,duplicados=salvar_itens_selecionados(editor)
                        st.success(f"{n} itens cadastrados; {duplicados} já existiam e foram ignorados.")
                        del st.session_state["importacao_previa"]
                        st.rerun()
                    except (ValueError,sqlite3.Error) as exc: st.error(str(exc))
                st.caption("A descrição/quantidade extraída de PDFs pode conter erros: confirme cada linha. Itens não marcados não são gravados.")
    with tabs[2]:
        with st.form("nova_ata"):
            pregao=st.text_input("Pregão (ex.: 35/2026) *")
            num=st.text_input("Número/ano da ata *")
            fornecedor=st.text_input("Fornecedor *")
            email=st.text_input("E-mail");telefone=st.text_input("Telefone")
            dt1=st.date_input("Início da vigência",value=date.today())
            dt2=st.date_input("Fim da vigência",value=date.today())
            url=st.text_input("Link do documento")
            if st.form_submit_button("Cadastrar"):
                try:
                    if not pregao.strip() or not num.strip() or not fornecedor.strip() or dt2<dt1: raise ValueError("Confira os campos obrigatórios e datas.")
                    execute("INSERT INTO atas(numero,pregao,fornecedor,email,telefone,inicio,fim,url) VALUES(?,?,?,?,?,?,?,?)",(num.strip(),pregao.strip(),fornecedor.strip(),email,telefone,dt1.isoformat(),dt2.isoformat(),url))
                    st.success("Cadastro realizado.")
                except (ValueError,sqlite3.Error) as exc:st.error(str(exc))
    with tabs[3]:
        atas=sql("SELECT * FROM atas ORDER BY pregao,numero");ps=sql("SELECT * FROM produtos ORDER BY nome")
        with st.form("novo_item_ata"):
            aid=pick("Ata (vinculada ao pregão)",atas,lambda r:f"{r['pregao']} · Ata {r['numero']}")
            pid=pick("Produto DEZOO",ps,lambda r:r["nome"])
            n=st.text_input("Número do item")
            qtd=st.number_input("Quantidade registrada",min_value=0.0)
            preco=st.number_input("Preço unitário R$",min_value=0.0,format="%.4f")
            unidade=st.text_input("Unidade do item",value="kg")
            if st.form_submit_button("Adicionar item") and aid and pid:
                try:
                    execute("INSERT INTO ata_itens(ata_id,produto_id,item_num,quantidade,preco,unidade) VALUES(?,?,?,?,?,?)",(aid,pid,n.strip(),qtd,preco,unidade.strip()))
                    st.success("Item cadastrado.")
                except sqlite3.IntegrityError:st.error("Item já cadastrado nesta ata.")


def pplan():
    st.title("Solicitações de recursos à PPLAN")
    st.caption("Autorização específica por item e quantidade; uma solicitação pode ter autorizações complementares.")
    a,b,c=st.tabs(["Nova solicitação","Atender solicitação","Histórico e memorando"])
    df=view_atas()
    with a:
        if df.empty: st.info("Cadastre ata e itens antes de solicitar recursos.")
        else:
            with st.form("pplan_nova"):
                numero=st.text_input("Número/referência do memorando")
                data=st.date_input("Data",key="pplan_dt")
                obs=st.text_area("Justificativa / observações")
                choices=options(df,lambda r:f"{r['produto']} · Ata {r['ata']} item {r['item_num']} · {money(r['preco'])}/{r['unidade']}")
                selected=st.multiselect("Itens solicitados",list(choices))
                qtds={}
                for s in selected:
                    qtds[choices[s]]=st.number_input(f"Quantidade · {s}",min_value=0.0,step=1.0,key=f"pplanqty_{choices[s]}")
                if st.form_submit_button("Salvar solicitação"):
                    if not selected or any(q<=0 for q in qtds.values()): st.error("Selecione itens com quantidades positivas.")
                    else:
                        pp=execute("INSERT INTO pplans (numero,data,situacao,observacoes) VALUES (?,?,?,?)",(numero,data.isoformat(),"Pendente",obs))
                        for ii,q in qtds.items(): execute("INSERT INTO pplan_itens (pplan_id,ata_item_id,quantidade) VALUES (?,?,?)",(pp,ii,q))
                        st.success("Solicitação registrada.")
    with b:
        rows=sql("""SELECT pi.id,p.nome AS produto,pp.numero,pi.quantidade,ai.unidade,COALESCE((SELECT SUM(quantidade) FROM autorizacoes au WHERE au.pplan_item_id=pi.id),0) AS autorizado
          FROM pplan_itens pi JOIN pplans pp ON pp.id=pi.pplan_id JOIN ata_itens ai ON ai.id=pi.ata_item_id JOIN produtos p ON p.id=ai.produto_id ORDER BY pp.id DESC""")
        table(rows)
        if not rows.empty:
            with st.form("autorizar"):
                selected=pick("Item a autorizar",rows,lambda r:f"{r['numero'] or 'Sem número'} · {r['produto']} · {r['autorizado']} / {r['quantidade']} {r['unidade']}")
                d=st.date_input("Data da autorização")
                q=st.number_input("Quantidade autorizada agora",min_value=0.0)
                ref=st.text_input("Referência do atendimento")
                if st.form_submit_button("Registrar autorização") and selected:
                    item=rows.set_index("id").loc[selected]
                    if q<=0 or q>float(item["quantidade"]-item["autorizado"])+1e-9: st.error("Quantidade inválida ou acima do restante solicitado.")
                    else:
                        execute("INSERT INTO autorizacoes (pplan_item_id,data,quantidade,referencia) VALUES (?,?,?,?)",(selected,d.isoformat(),q,ref))
                        st.success("Autorização registrada.")
    with c:
        table(sql("SELECT id,numero,data,situacao,observacoes FROM pplans ORDER BY id DESC"))
        pp=sql("SELECT * FROM pplans ORDER BY id DESC")
        if not pp.empty:
            ppid=pick("Memorando de solicitação",pp,lambda r:r["numero"] or f"Solicitação {r['id']}")
            if ppid:
                linhas=sql("""SELECT p.nome,pi.quantidade,ai.unidade,ai.preco,a.numero AS ata FROM pplan_itens pi JOIN ata_itens ai ON ai.id=pi.ata_item_id JOIN atas a ON a.id=ai.ata_id JOIN produtos p ON p.id=ai.produto_id WHERE pi.pplan_id=?""",(ppid,))
                rem=st.text_input("Remetente",value="Caio Rodrigues Monteiro")
                just=st.text_area("Texto da justificativa",value="Solicitamos a autorização de recursos para aquisição dos materiais relacionados, destinados às atividades do Departamento de Zootecnia.")
                d=Document();d.add_heading("Solicitação de recursos à PPLAN",0);d.add_paragraph(f"Remetente: {rem}");d.add_paragraph(just)
                t=d.add_table(rows=1, cols=5);t.style="Table Grid"
                for cell,v in zip(t.rows[0].cells,["Produto","Ata","Quantidade","Unidade","Valor (R$)"]): cell.text=v
                for _,r in linhas.iterrows():
                    for cell,v in zip(t.add_row().cells,[r["nome"],r["ata"],str(r["quantidade"]),r["unidade"],money(r["quantidade"]*r["preco"])]): cell.text=v
                bfr=io.BytesIO();d.save(bfr)
                st.download_button("Baixar rascunho Word do memorando",bfr.getvalue(),"memorando_pplan_rascunho.docx")
                st.info("Rascunho genérico: falta incorporar o modelo oficial da UFSJ.")


def compras():
    st.title("Requisições SIPAC e empenhos")
    tab1,tab2,tab3,tab4=st.tabs(["Nova requisição","Novo empenho","Itens empenhados","Consultas"])
    items=view_atas()
    reqs=sql("SELECT * FROM requisicoes ORDER BY id DESC")
    with tab1:
        st.caption("A requisição SIPAC pode ser cadastrada por item. O vínculo PPLAN é opcional para migrar registros antigos.")
        with st.form("requisicao_form"):
            num=st.text_input("Número/ano da requisição *"); dt=st.date_input("Data",key="req_dt")
            iid=pick("Item da ata",items,lambda r:f"{r['produto']} · {r['ata']} item {r['item_num']}")
            q=st.number_input("Quantidade requisitada",min_value=0.0)
            pi_df=sql("SELECT pi.id,pp.numero,p.nome AS produto FROM pplan_itens pi JOIN pplans pp ON pp.id=pi.pplan_id JOIN ata_itens ai ON ai.id=pi.ata_item_id JOIN produtos p ON p.id=ai.produto_id ORDER BY pi.id DESC")
            p_opts={"Sem vínculo (histórico)":None,**options(pi_df,lambda r:f"{r['numero']} · {r['produto']}")}
            p_sel=st.selectbox("Solicitação PPLAN vinculada",list(p_opts))
            if st.form_submit_button("Registrar requisição e item"):
                if not num.strip() or q<=0 or iid is None: st.error("Preencha número, item e quantidade.")
                else:
                    rid=execute("INSERT OR IGNORE INTO requisicoes (numero,data) VALUES (?,?)",(num.strip(),dt.isoformat()))
                    if rid==0: rid=one("SELECT id FROM requisicoes WHERE numero=?",(num.strip(),))["id"]
                    execute("INSERT INTO req_itens (requisicao_id,ata_item_id,pplan_item_id,quantidade) VALUES (?,?,?,?)",(rid,iid,p_opts[p_sel],q))
                    st.success("Item da requisição registrado.")
    with tab2:
        with st.form("emp_novo"):
            num=st.text_input("Número/ano do empenho *")
            opts={"Sem requisição vinculada":None,**options(reqs,lambda r:r["numero"])}
            s=st.selectbox("Requisição de origem",list(opts))
            dt=st.date_input("Data de emissão")
            if st.form_submit_button("Cadastrar empenho"):
                try:
                    if not num.strip(): raise ValueError("Informe o número do empenho.")
                    execute("INSERT INTO empenhos (numero,requisicao_id,data) VALUES (?,?,?)",(num.strip(),opts[s],dt.isoformat()))
                    st.success("Empenho cadastrado.")
                except (ValueError,sqlite3.IntegrityError) as exc: st.error(str(exc))
    with tab3:
        emps=sql("SELECT * FROM empenhos ORDER BY id DESC")
        with st.form("emp_item"):
            eid=pick("Empenho",emps,lambda r:r["numero"])
            iid=pick("Produto da ata",items,lambda r:f"{r['produto']} · {r['ata']} item {r['item_num']}")
            q=st.number_input("Quantidade empenhada",min_value=0.0)
            preco=st.number_input("Valor unitário empenhado R$",min_value=0.0,format="%.4f")
            if st.form_submit_button("Adicionar item ao empenho"):
                if eid and iid and q>0:
                    execute("INSERT INTO empenho_itens (empenho_id,ata_item_id,quantidade,preco) VALUES (?,?,?,?)",(eid,iid,q,preco))
                    st.success("Item empenhado registrado.")
                else: st.error("Selecione empenho, item e quantidade positiva.")
    with tab4:
        table(sql("SELECT r.numero AS requisicao,ai.item_num,p.nome AS produto,ri.quantidade,ai.unidade FROM req_itens ri JOIN requisicoes r ON r.id=ri.requisicao_id JOIN ata_itens ai ON ai.id=ri.ata_item_id JOIN produtos p ON p.id=ai.produto_id ORDER BY r.id DESC"))
        table(pendencias())
        st.info("Consulta automática a APIs do governo ainda não habilitada.")


def recebimentos():
    st.title("Recebimentos, DANFEs e SEALM")
    st.caption("Vários DANFEs por entrega e repetição do mesmo item em notas diferentes. Quantidades acumuladas por empenho.")
    x,y,z,w=st.tabs(["1. Registrar entrega","2. Incluir DANFE","3. Encaminhar ao SEALM","4. Acumulado por empenho"])
    emps=sql("SELECT * FROM empenhos ORDER BY id DESC")
    with x:
        with st.form("entrega"):
            eid=pick("Empenho",emps,lambda r:r["numero"])
            dt=st.date_input("Data da entrega")
            local=st.text_input("Local",value="Campus Tancredo Neves (CTAN)")
            conf=st.checkbox("Produtos conferidos: quantidade e qualidade",value=False)
            obs=st.text_area("Observações")
            if st.form_submit_button("Cadastrar entrega") and eid:
                execute("INSERT INTO entregas (empenho_id,data,local,conferido,observacoes) VALUES (?,?,?,?,?)",(eid,dt.isoformat(),local,int(conf),obs))
                st.success("Entrega registrada. Agora cadastre os DANFEs e distribua os itens.")
    with y:
        entregas=sql("SELECT ent.id,ent.data,e.numero AS empenho FROM entregas ent JOIN empenhos e ON e.id=ent.empenho_id ORDER BY ent.id DESC")
        with st.form("nf_form"):
            enid=pick("Entrega",entregas,lambda r:f"{r['data']} · {r['empenho']}")
            n=st.text_input("Número do DANFE *")
            serie=st.text_input("Série",value="1")
            val=st.number_input("Valor total da nota R$ (opcional)",min_value=0.0)
            if st.form_submit_button("Cadastrar DANFE"):
                if enid and n.strip():
                    existing=one("SELECT id FROM notas WHERE entrega_id=? AND numero=? AND serie=?",(enid,n.strip(),serie.strip()))
                    if existing: st.error("Esta nota já está registrada nesta entrega.")
                    else:
                        execute("INSERT INTO notas (entrega_id,numero,serie,valor) VALUES (?,?,?,?)",(enid,n.strip(),serie.strip(),val))
                        st.success("DANFE cadastrado. Relacione agora seus produtos.")
                else: st.error("Selecione entrega e número da nota.")
        notas=sql("""SELECT n.id,n.numero,n.serie,n.entrega_id,e.numero AS empenho,ent.data,ent.empenho_id FROM notas n JOIN entregas ent ON ent.id=n.entrega_id JOIN empenhos e ON e.id=ent.empenho_id ORDER BY n.id DESC""")
        if not notas.empty:
            nf=pick("DANFE para incluir itens",notas,lambda r:f"{r['numero']} série {r['serie']} · {r['empenho']}")
            if nf:
                empenho_id=int(notas.set_index("id").loc[nf,"empenho_id"])
                eitems=sql("SELECT ei.id,p.nome AS produto,ei.quantidade,ai.unidade FROM empenho_itens ei JOIN ata_itens ai ON ai.id=ei.ata_item_id JOIN produtos p ON p.id=ai.produto_id WHERE ei.empenho_id=?",(empenho_id,))
                with st.form("nf_item_form"):
                    item_id=pick("Item empenhado",eitems,lambda r:f"{r['produto']} · {r['quantidade']} {r['unidade']}")
                    qtd=st.number_input("Quantidade nesta nota",min_value=0.0)
                    if st.form_submit_button("Adicionar produto ao DANFE") and item_id:
                        p=pendencias()
                        pending=float(p.loc[p.empenho_item_id==item_id,"pendente"].iloc[0]) if not p.loc[p.empenho_item_id==item_id].empty else 0
                        # Include all recorded notes, even unchecked deliveries, to avoid duplicate physical allocations.
                        registered=one("SELECT COALESCE(SUM(quantidade),0) AS q FROM nota_itens WHERE empenho_item_id=?",(item_id,))["q"]
                        planned=float(eitems.loc[eitems.id==item_id,"quantidade"].iloc[0])
                        if qtd<=0 or qtd>planned-float(registered)+1e-9: st.error("Quantidade excede o saldo ainda não distribuído entre DANFEs.")
                        else:
                            execute("INSERT INTO nota_itens (nota_id,empenho_item_id,quantidade) VALUES (?,?,?)",(nf,item_id,qtd))
                            st.success("Produto alocado ao DANFE sem duplicar a entrega.")
                table(sql("SELECT p.nome AS produto,ni.quantidade,ai.unidade FROM nota_itens ni JOIN empenho_itens ei ON ei.id=ni.empenho_item_id JOIN ata_itens ai ON ai.id=ei.ata_item_id JOIN produtos p ON p.id=ai.produto_id WHERE ni.nota_id=?",(nf,)))
    with z:
        nfs=sql("SELECT n.id,n.numero,n.serie,n.enviado,n.data_envio,ent.data AS recebimento,ent.conferido,e.numero AS empenho FROM notas n JOIN entregas ent ON ent.id=n.entrega_id JOIN empenhos e ON e.id=ent.empenho_id ORDER BY n.id DESC")
        table(nfs)
        pend=nfs[(nfs.enviado==0)&(nfs.conferido==1)] if not nfs.empty else nfs
        if not pend.empty:
            options_nf=options(pend,lambda r:f"DANFE {r['numero']} · {r['empenho']} · {r['recebimento']}")
            selected=st.multiselect("Selecionar DANFEs para um e-mail",list(options_nf))
            nome=st.text_input("Nome do remetente",value="Caio Rodrigues Monteiro")
            cargo=st.text_input("Cargo / setor",value="Técnico em Agropecuária - DEZOO/UFSJ")
            if selected:
                ids=[options_nf[k] for k in selected]
                placeholders=",".join("?"*len(ids))
                linhas=sql(f"SELECT n.id,n.numero,e.numero AS empenho,ent.data, p.nome AS produto,SUM(ni.quantidade) AS quantidade,ai.unidade FROM notas n JOIN entregas ent ON ent.id=n.entrega_id JOIN empenhos e ON e.id=ent.empenho_id LEFT JOIN nota_itens ni ON ni.nota_id=n.id LEFT JOIN empenho_itens ei ON ei.id=ni.empenho_item_id LEFT JOIN ata_itens ai ON ai.id=ei.ata_item_id LEFT JOIN produtos p ON p.id=ai.produto_id WHERE n.id IN ({placeholders}) GROUP BY n.id,n.numero,e.numero,ent.data,p.nome,ai.unidade",ids)
                itens="\n".join(f"- {r['produto']}: {r['quantidade']:g} {r['unidade']} (DANFE {r['numero']})" for _,r in linhas.dropna(subset=['produto']).iterrows())
                notas_txt=", ".join(pend[pend.id.isin(ids)].numero.astype(str).tolist())
                empenhos_txt=", ".join(sorted(set(pend[pend.id.isin(ids)].empenho.astype(str))))
                texto=f"Prezados, bom dia.\n\nCom os cordiais cumprimentos, venho, por meio deste, atestar o recebimento dos produtos relacionados:\n\n{itens}\n\nConforme os DANFEs {notas_txt}, que deverão ser anexados à mensagem.\n\nEmpenho(s): {empenhos_txt}. Atesto a adequada qualidade dos produtos e o cumprimento das condições acordadas, conforme conferência registrada. Encaminho as notas para providências relativas ao pagamento.\n\nAtenciosamente,\n{nome}\n{cargo}"
                st.text_area("Rascunho de e-mail para SEALM",value=texto,height=300)
                st.download_button("Baixar e-mail em TXT",texto,"ateste_sealm.txt")
                if st.button("Confirmar que encaminhei esses DANFEs ao SEALM"):
                    for _id in ids: execute("UPDATE notas SET enviado=1,data_envio=? WHERE id=?",(date.today().isoformat(),_id))
                    st.success("Envio informado manualmente. Nenhum e-mail foi enviado pelo aplicativo.")
        else: st.info("Sem notas pendentes com entrega conferida.")
    with w:
        df=pendencias()
        table(df[["empenho","produto","fornecedor","empenhado","recebido","pendente"]] if not df.empty else df)
        st.caption("Somente notas vinculadas a entregas marcadas como conferidas entram no acumulado confirmado.")


def orcamentos():
    st.title("Pesquisa de preços")
    st.caption("Geração de rascunhos. Cotação e contato não são enviados automaticamente.")
    nome=st.text_input("Nome de quem envia",value="Caio Rodrigues Monteiro",key="orc_nome")
    cargo=st.text_input("Cargo",value="Técnico em Agropecuária",key="orc_cargo")
    destinatario=st.text_input("Empresa destinatária (opcional)")
    grupos=st.text_area("Grupos de produtos",value="Alimentos para animais e gêneros de alimentação (3006 e 3007); Animais para pesquisa e abate (3008)")
    email=f"Bom dia,\n\nMeu nome é {nome} e sou {cargo} do Departamento de Zootecnia da Universidade Federal de São João del-Rei (UFSJ).\n\nEstamos organizando um processo de registro de preços. Para subsidiar a pesquisa de preços, solicitamos orçamento dos itens relacionados nos modelos anexos.\n\nGrupos: {grupos}.\n\nPreferencialmente, solicitamos preenchimento dos modelos encaminhados, contemplando todas as informações solicitadas.\n\nDesde já, agradecemos a atenção.\n\nAtenciosamente,\n{nome}\n{cargo} · DEZOO/UFSJ"
    st.text_area("Rascunho editável para copiar",value=email,height=280)
    st.download_button("Baixar rascunho",email,"solicitacao_orcamento.txt")
    st.info("Na versão posterior, serão incluídos cotações por empresa, acompanhamento de respostas, anexos e comparativo de preços.")


def relatorios():
    st.title("Relatórios prioritários")
    a,b,c=st.tabs(["Entregas pendentes","Saldo dos itens de atas","Notas pendentes de envio"])
    with a:
        df=pendencias()
        if not df.empty:
            df=df[df.pendente>0.000001].copy()
            df["valor_pendente"]=df.pendente*df.preco
        table(df.drop(columns=["empenho_item_id","preco"],errors="ignore"))
        download_xlsx(df.drop(columns=["empenho_item_id"],errors="ignore"),"dzoo_entregas_pendentes.xlsx")
    with b:
        df=view_atas()
        if not df.empty:
            df["saldo_calculado"]=df.quantidade-df.empenhado
            df=df[df.saldo_calculado>0].copy()
            df["valor_potencial"]=df.saldo_calculado*df.preco
        st.warning("Saldo interno = quantidade registrada – quantidade empenhada cadastrada. Não é consulta oficial ao SIPAC.")
        table(df.drop(columns=["id","ata_id"],errors="ignore"))
        download_xlsx(df.drop(columns=["id","ata_id"],errors="ignore"),"dzoo_saldo_interno_atas.xlsx")
    with c:
        df=sql("SELECT n.numero AS danfe,n.serie,e.numero AS empenho,ent.data AS entrega,a.fornecedor FROM notas n JOIN entregas ent ON ent.id=n.entrega_id JOIN empenhos e ON e.id=ent.empenho_id JOIN nota_itens ni ON ni.nota_id=n.id JOIN empenho_itens ei ON ei.id=ni.empenho_item_id JOIN ata_itens ai ON ai.id=ei.ata_item_id JOIN atas a ON a.id=ai.ata_id WHERE n.enviado=0 GROUP BY n.id")
        table(df); download_xlsx(df,"dzoo_danfes_pendentes.xlsx")


def sidebar():
    with st.sidebar:
        st.markdown("### 🐑 DZOO · UFSJ")
        menu=st.radio("Navegar",["Painel inicial","Produtos e setores","Planejamento anual","Licitações e atas","PPLAN","Requisições e empenhos","Recebimentos e SEALM","Pesquisa de preços","Relatórios"],label_visibility="collapsed")
        st.divider()
        st.warning("**PROTÓTIPO LOCAL**\n\nSem login, sem integrações oficiais, sem envio de e-mails.")
        if st.button("Popular dados fictícios",use_container_width=True):
            seed(); st.rerun()
        if st.button("Limpar dados de demonstração",use_container_width=True):
            wipe_demo(); st.rerun()
    return menu


init()
choice=sidebar()
{
    "Painel inicial":home,
    "Produtos e setores":produtos,
    "Planejamento anual":planejamento,
    "Licitações e atas":licitacoes,
    "PPLAN":pplan,
    "Requisições e empenhos":compras,
    "Recebimentos e SEALM":recebimentos,
    "Pesquisa de preços":orcamentos,
    "Relatórios":relatorios,
}[choice]()
